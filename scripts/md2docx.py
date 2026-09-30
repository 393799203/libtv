#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把讲书稿 Markdown 转成排版规范的 Word 文档（中文字体 + 真表格 + 页码）。

用法: python3 scripts/md2docx.py <输入.md> <输出.docx>
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CJK_BODY = '宋体'
CJK_HEAD = '微软雅黑'
LATIN = 'Times New Roman'

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def set_run(run, size=11, bold=False, color=None, cjk=CJK_BODY, latin=LATIN):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = latin
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), latin)
    rf.set(qn('w:hAnsi'), latin)
    rf.set(qn('w:eastAsia'), cjk)


def add_runs(par, text, size=11, base_bold=False, color=None, cjk=CJK_BODY):
    """处理 **加粗** 与 `代码` 两种行内标记"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            set_run(par.add_run(part[2:-2]), size=size, bold=True, color=color, cjk=cjk)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = par.add_run(part[1:-1])
            set_run(r, size=size - 0.5, bold=base_bold, color=RGBColor(0xB0, 0x30, 0x60), latin='Consolas')
        else:
            set_run(par.add_run(part), size=size, bold=base_bold, color=color, cjk=cjk)


def shade(el, hex_fill):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hex_fill)
    el.append(sh)


def left_bar_and_shade(par, fill='F4F7FB', bar='4472C4'):
    """给段落加左侧竖条 + 底纹（口播稿专用）"""
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), bar)
    pbdr.append(left)
    ppr.append(pbdr)
    shade(ppr, fill)


def add_hr(par):
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'BFBFBF')
    pbdr.append(bottom)
    ppr.append(pbdr)


def setup_page(doc):
    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.4)
        s.top_margin = s.bottom_margin = Cm(2.2)


def add_page_number_footer(doc):
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        set_run(r, size=9, color=RGBColor(0x80, 0x80, 0x80))
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        r._element.addnext(fld)


def style_heading(doc, style_name, size, color):
    st = doc.styles[style_name]
    st.font.name = LATIN
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK_HEAD)
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.3


def style_doc_defaults(doc):
    st = doc.styles['Normal']
    st.font.name = LATIN
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK_BODY)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(4)


def build_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, cell_text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ''
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.15
        par.paragraph_format.space_after = Pt(2)
        add_runs(par, cell_text, size=10, base_bold=True)
        shade(cell._tc.get_or_add_tcPr(), 'DCE6F1')
    for row in body:
        cells = t.add_row().cells
        for i, cell_text in enumerate(row[:len(header)]):
            cells[i].text = ''
            par = cells[i].paragraphs[0]
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, cell_text, size=10)
    doc.add_paragraph()


def split_table_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def convert(src, dst):
    lines = open(src, encoding='utf-8').read().split('\n')
    doc = Document()
    setup_page(doc)
    style_doc_defaults(doc)
    style_heading(doc, 'Title', 20, RGBColor(0x1F, 0x38, 0x64))
    style_heading(doc, 'Heading 1', 15, RGBColor(0x1F, 0x38, 0x64))
    style_heading(doc, 'Heading 2', 12.5, RGBColor(0x1F, 0x38, 0x64))
    add_page_number_footer(doc)

    i = 0
    seen_quote = False
    page_break_done = False  # 只在顶部元信息块后分一次页
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # 表格
        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            rows = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_table_row(lines[i]))
                i += 1
            build_table(doc, rows)
            continue

        # 分隔线
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            add_hr(p)
            i += 1
            continue

        # 标题
        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            text = stripped[level:].strip()
            if level == 1:
                p = doc.add_paragraph(style='Title')
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                add_runs(p, text, size=20, base_bold=True, cjk=CJK_HEAD)
            else:
                p = doc.add_paragraph(style='Heading 1' if level == 2 else 'Heading 2')
                size = 15 if level == 2 else 12.5
                add_runs(p, text, size=size, base_bold=True,
                         color=RGBColor(0x1F, 0x38, 0x64), cjk=CJK_HEAD)
            i += 1
            continue

        # 引用（口播稿）
        if stripped.startswith('>'):
            text = stripped.lstrip('>').strip()
            if not text:
                i += 1
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.right_indent = Cm(0.3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            left_bar_and_shade(p)
            add_runs(p, text, size=11.5)
            seen_quote = True
            i += 1
            continue

        # 空行
        if not stripped:
            # 元信息块结束后分页一次，正文从第 2 页开始
            if seen_quote and not page_break_done and i > 0 and lines[i - 1].strip().startswith('>'):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                page_break_done = True
            i += 1
            continue

        # 【画面】【口播】【提示】等操作标记
        m = re.match(r'^【(.{1,4})】(.*)$', stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            tag = m.group(1)
            color = {'画面': RGBColor(0x1F, 0x6F, 0x3C), '提示': RGBColor(0xB0, 0x60, 0x00)}.get(tag, RGBColor(0x1F, 0x38, 0x64))
            set_run(p.add_run(f'【{tag}】'), size=11, bold=True, color=color, cjk=CJK_HEAD)
            add_runs(p, m.group(2), size=11)
            i += 1
            continue

        # 有序列表：保留原编号，避免 Word 自动编号串号
        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            set_run(p.add_run(f'{m.group(1)}. '), size=11, bold=True)
            add_runs(p, m.group(2), size=11)
            i += 1
            continue

        # 无序列表
        if re.match(r'^[-*]\s+', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            set_run(p.add_run('· '), size=11, bold=True)
            add_runs(p, re.sub(r'^[-*]\s+', '', stripped), size=11)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped, size=11)
        i += 1

    doc.save(dst)
    return dst


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    out = convert(sys.argv[1], sys.argv[2])
    print(f'✅ 已生成 {out}')